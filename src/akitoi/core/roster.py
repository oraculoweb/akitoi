"""
Roster (padrón) bulk import/export for clubs.

The admin manages the club as ONE file: upload a CSV/XLSX with every
member and their state, re-upload it whenever the roster changes.

Template columns (header row required, in any order):

    nombre,email,telefono,rol,estado,valid_until

- nombre      required
- email       required — the UNIQUE UPSERT KEY within the club
- telefono    optional
- rol         optional (default "socio")
- estado      vigente | suspendido | vencido | baja  (default vigente)
- valid_until optional, ISO (2027-12-31) or DD/MM/YYYY

Import semantics (upsert by email):
- unknown email            -> member created with their state
- known email, new state   -> updated in place (mass revocation: a row
                              flipped to "suspendido" suspends the card
                              on the next scan)
- estado "baja"            -> SOFT DELETE: the record persists for
                              audit (Ley 29733) but leaves active views
                              and the card dies instantly
- malformed row            -> reported in the import log, the rest of
                              the file still imports (partial import)

Every import is recorded (timestamp + created/updated/deactivated/
error counts) in a queryable import log.
"""
import csv
import io
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from typing import List, Optional, Tuple

from ..models.member import Member
from ..models.membership import Membership, MembershipStatus
from ..storage.org_base import OrgStorageBackend
from ..utils.validators import validate_email

ROSTER_COLUMNS = ["nombre", "email", "telefono", "rol", "estado", "valid_until"]
VALID_STATES = {"vigente", "suspendido", "vencido", "baja"}

ROSTER_TEMPLATE_CSV = (
    "nombre,email,telefono,rol,estado,valid_until\r\n"
    "Ana Pérez,ana@example.com,+51911222333,socio,vigente,2027-12-31\r\n"
    "Luis Gómez,luis@example.com,,directivo,suspendido,2027-12-31\r\n"
    "Rosa Díaz,rosa@example.com,+51944555666,socio,baja,\r\n"
)


@dataclass
class ImportLog:
    """Auditable record of one roster import."""

    organization_id: str
    created: int = 0
    updated: int = 0
    deactivated: int = 0
    errors: List[dict] = field(default_factory=list)
    filename: Optional[str] = None
    id: str = field(default_factory=lambda: str(uuid.uuid4()))
    imported_at: datetime = field(default_factory=datetime.now)

    @property
    def error_count(self) -> int:
        return len(self.errors)

    def to_dict(self) -> dict:
        return {
            "id": self.id,
            "organization_id": self.organization_id,
            "imported_at": self.imported_at.isoformat(),
            "created": self.created,
            "updated": self.updated,
            "deactivated": self.deactivated,
            "error_count": self.error_count,
            "errors": self.errors,
            "filename": self.filename,
        }


# --- File parsing ---


def parse_roster_file(filename: str, content: bytes) -> List[dict]:
    """
    Parse a roster file into raw rows (dicts keyed by template columns).

    Args:
        filename: Original filename (extension picks the parser)
        content: Raw file bytes

    Returns:
        List of {"_row": <line number>, <column>: <value>, ...}

    Raises:
        ValueError: Unsupported extension or missing required columns
    """
    lower = filename.lower()
    if lower.endswith(".csv"):
        rows = _parse_csv(content)
    elif lower.endswith((".xlsx", ".xls")):
        rows = _parse_xlsx(content)
    else:
        raise ValueError(
            f"Formato no soportado: {filename}. Usa .csv o .xlsx"
        )
    return rows


def _normalize_header(header: List[str]) -> List[str]:
    normalized = [str(h or "").strip().lower() for h in header]
    missing = [c for c in ROSTER_COLUMNS if c not in normalized]
    if missing:
        raise ValueError(f"Columnas faltantes en el archivo: {', '.join(missing)}")
    return normalized


def _parse_csv(content: bytes) -> List[dict]:
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    try:
        header = _normalize_header(next(reader))
    except StopIteration:
        raise ValueError("Archivo vacío")

    rows = []
    for line_number, values in enumerate(reader, start=2):
        if not any(str(v).strip() for v in values):
            continue  # skip blank lines
        row = dict(zip(header, [str(v).strip() for v in values]))
        row["_row"] = line_number
        rows.append(row)
    return rows


def _parse_xlsx(content: bytes) -> List[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = _normalize_header(list(next(iterator)))
    except StopIteration:
        raise ValueError("Archivo vacío")

    rows = []
    for line_number, values in enumerate(iterator, start=2):
        cells = ["" if v is None else v for v in values]
        if not any(str(v).strip() for v in cells):
            continue
        row = {}
        for key, value in zip(header, cells):
            if isinstance(value, datetime):
                row[key] = value.date().isoformat()
            else:
                row[key] = str(value).strip()
        row["_row"] = line_number
        rows.append(row)
    workbook.close()
    return rows


# --- Row validation ---


def _parse_valid_until(value: str) -> Optional[datetime]:
    if not value:
        return None
    for parser in (
        lambda v: datetime.fromisoformat(v),
        lambda v: datetime.strptime(v, "%d/%m/%Y"),
    ):
        try:
            parsed = parser(value)
            # Dates without time mean "valid through that whole day"
            if parsed.hour == 0 and parsed.minute == 0 and parsed.second == 0:
                return parsed.replace(hour=23, minute=59, second=59)
            return parsed
        except ValueError:
            continue
    raise ValueError(f"Fecha inválida: {value} (usa YYYY-MM-DD o DD/MM/YYYY)")


def _clean_row(row: dict) -> dict:
    """Validate one raw row. Raises ValueError with a readable reason."""
    name = row.get("nombre", "").strip()
    if not name:
        raise ValueError("nombre es obligatorio")

    email = row.get("email", "").strip().lower()
    if not email:
        raise ValueError("email es obligatorio (es la clave de upsert)")
    if not validate_email(email):
        raise ValueError(f"email inválido: {email}")

    estado = (row.get("estado", "") or "vigente").strip().lower()
    if estado not in VALID_STATES:
        raise ValueError(
            f"estado inválido: {estado} (usa vigente|suspendido|vencido|baja)"
        )

    return {
        "name": name,
        "email": email,
        "phone": row.get("telefono", "").strip() or None,
        "role": row.get("rol", "").strip() or "socio",
        "estado": estado,
        "valid_until": _parse_valid_until(row.get("valid_until", "").strip()),
    }


# --- Import (upsert) ---


def import_roster(
    storage: OrgStorageBackend,
    organization_id: str,
    rows: List[dict],
    filename: Optional[str] = None,
) -> ImportLog:
    """
    Upsert a parsed roster into an organization.

    Malformed rows are collected into the log without aborting the
    rest of the file. The resulting ImportLog is persisted.
    """
    log = ImportLog(organization_id=organization_id, filename=filename)
    existing_by_email = {
        member.email.strip().lower(): member
        for member in storage.list_members(organization_id)
        if member.email
    }

    for row in rows:
        try:
            data = _clean_row(row)
        except ValueError as exc:
            log.errors.append({"row": row.get("_row"), "error": str(exc)})
            continue

        member = existing_by_email.get(data["email"])
        if member is None:
            _create_from_row(storage, organization_id, data, log)
        else:
            _update_from_row(storage, member, data, log)

    storage.save_import_log(log.to_dict())
    return log


def _create_from_row(
    storage: OrgStorageBackend, organization_id: str, data: dict, log: ImportLog
) -> None:
    member = Member(
        organization_id=organization_id,
        name=data["name"],
        email=data["email"],
        phone=data["phone"],
        role=data["role"],
    )
    membership = Membership(
        member_id=member.id,
        organization_id=organization_id,
        valid_until=data["valid_until"],
    )
    if data["estado"] == "baja":
        member.deactivate()
        membership.expire()
        log.deactivated += 1
    else:
        membership.status = MembershipStatus(data["estado"])
        log.created += 1
    storage.save_member(member)
    storage.save_membership(membership)


def _update_from_row(
    storage: OrgStorageBackend, member: Member, data: dict, log: ImportLog
) -> None:
    member.name = data["name"]
    member.phone = data["phone"]
    member.role = data["role"]
    member.updated_at = datetime.now()

    membership = storage.get_membership_by_member(member.id)
    if membership is None:
        membership = Membership(
            member_id=member.id, organization_id=member.organization_id
        )

    if data["estado"] == "baja":
        if member.is_active:
            member.deactivate()
        membership.expire()
        log.deactivated += 1
    else:
        if not member.is_active:
            member.restore()  # re-alta via roster
        membership.status = MembershipStatus(data["estado"])
        membership.valid_until = data["valid_until"]
        membership.updated_at = datetime.now()
        log.updated += 1

    storage.save_member(member)
    storage.save_membership(membership)


# --- Export (report) ---


def member_estado(member: Member, membership: Optional[Membership]) -> str:
    """Effective state as the roster report should show it."""
    if not member.is_active:
        return "baja"
    if membership is None:
        return "vigente"
    if (
        membership.status == MembershipStatus.VIGENTE
        and not membership.is_active()
    ):
        return "vencido"
    return membership.status.value


def _report_rows(
    storage: OrgStorageBackend, organization_id: str
) -> List[List[str]]:
    rows = []
    members = sorted(
        storage.list_members(organization_id), key=lambda m: m.name.lower()
    )
    for member in members:
        membership = storage.get_membership_by_member(member.id)
        valid_until = (
            membership.valid_until.date().isoformat()
            if membership and membership.valid_until
            else ""
        )
        last_update = max(
            [member.updated_at]
            + ([membership.updated_at] if membership else [])
        )
        rows.append([
            member.name,
            member.email or "",
            member.phone or "",
            member.role,
            member_estado(member, membership),
            valid_until,
            last_update.strftime("%Y-%m-%d %H:%M"),
        ])
    return rows


EXPORT_COLUMNS = ROSTER_COLUMNS + ["actualizado"]


def export_roster_csv(storage: OrgStorageBackend, organization_id: str) -> bytes:
    """Export the current roster (including bajas) as CSV bytes."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_COLUMNS)
    writer.writerows(_report_rows(storage, organization_id))
    return buffer.getvalue().encode("utf-8-sig")


def export_roster_xlsx(storage: OrgStorageBackend, organization_id: str) -> bytes:
    """Export the current roster (including bajas) as XLSX bytes."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Padrón"
    sheet.append(EXPORT_COLUMNS)
    for row in _report_rows(storage, organization_id):
        sheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
